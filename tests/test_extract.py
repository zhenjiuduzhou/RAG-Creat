import importlib.util
import sys
from pathlib import Path

import pytest

from ragstudio.extract import ExtractionError, Section, extract_file


def test_text_file_returns_document_section(tmp_path: Path) -> None:
    path = tmp_path / "note.TXT"
    path.write_text("  Hello extraction.  ", encoding="utf-8")

    result = extract_file(path)

    assert result == [Section("Hello extraction.", "document", "text")]


def test_csv_decodes_gb18030_and_keeps_row_locations(tmp_path: Path) -> None:
    path = tmp_path / "table.csv"
    path.write_bytes("名称,数量\n棉花,3\n".encode("gb18030"))

    result = extract_file(path)

    assert [section.text for section in result] == ["名称\t数量", "名称\t数量\n棉花\t3"]
    assert [section.location for section in result] == ["row 1", "row 2"]


def test_empty_and_unsupported_files_raise_clear_errors(tmp_path: Path) -> None:
    empty = tmp_path / "empty.txt"
    empty.touch()
    unsupported = tmp_path / "archive.zip"
    unsupported.write_bytes(b"zip")

    with pytest.raises(ExtractionError, match="empty"):
        extract_file(empty)
    with pytest.raises(ExtractionError, match="Unsupported"):
        extract_file(unsupported)


def test_missing_file_and_oversize_file_are_rejected(tmp_path: Path) -> None:
    with pytest.raises(ExtractionError, match="does not exist"):
        extract_file(tmp_path / "missing.txt")

    large = tmp_path / "large.txt"
    with large.open("wb") as file:
        file.truncate(100 * 1024 * 1024 + 1)
    with pytest.raises(ExtractionError, match="100 MB"):
        extract_file(large)


def test_invalid_json_is_reported_as_extraction_error(tmp_path: Path) -> None:
    path = tmp_path / "bad.json"
    path.write_text("{not valid", encoding="utf-8")
    with pytest.raises(ExtractionError, match="Cannot extract JSON"):
        extract_file(path)


@pytest.mark.skipif(importlib.util.find_spec("docx") is None, reason="python-docx is unavailable")
def test_docx_paragraph_locations(tmp_path: Path) -> None:
    from docx import Document

    path = tmp_path / "sample.docx"
    document = Document()
    document.add_paragraph("First paragraph")
    document.add_paragraph("")
    document.add_paragraph("Third paragraph")
    document.save(path)

    result = extract_file(path)

    assert [section.text for section in result] == ["First paragraph", "Third paragraph"]
    assert [section.location for section in result] == ["paragraph 1", "paragraph 3"]


def test_docx_heading_and_table_headers_stay_with_answers(tmp_path: Path) -> None:
    from docx import Document

    path = tmp_path / "policy.docx"
    document = Document()
    document.add_heading("产品 A 质保期限", level=1)
    document.add_paragraph("两年")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "产品"
    table.cell(0, 1).text = "金额"
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "100 元"
    document.save(path)

    result = extract_file(path)
    assert "产品 A 质保期限\n两年" in [section.text for section in result]
    assert "产品\t金额\nA\t100 元" in [section.text for section in result]


@pytest.mark.skipif(sys.platform != "win32" or importlib.util.find_spec("rapidocr") is None,
                    reason="Windows OCR dependencies are unavailable")
def test_chinese_image_and_scanned_pdf_ocr(tmp_path: Path) -> None:
    from PIL import Image, ImageDraw, ImageFont

    font_path = Path(r"C:\Windows\Fonts\msyh.ttc")
    if not font_path.exists():
        pytest.skip("Chinese font is unavailable")
    image = Image.new("RGB", (900, 180), "white")
    ImageDraw.Draw(image).text((25, 45), "企业知识库 测试 ABC 123",
                               font=ImageFont.truetype(str(font_path), 48), fill="black")
    image_path = tmp_path / "scan.png"
    image.save(image_path)
    assert "企业知识库" in "".join(row.text for row in extract_file(image_path))

    pdf_path = tmp_path / "scan.pdf"
    image.save(pdf_path, "PDF")
    sections = extract_file(pdf_path)
    assert sections[0].location == "page 1"
    assert sections[0].method == "ocr"
    assert "企业知识库" in sections[0].text


def test_office_web_and_json_formats(tmp_path: Path) -> None:
    from openpyxl import Workbook
    from pptx import Presentation

    sheet_file = tmp_path / "table.xlsx"
    book = Workbook()
    book.active.title = "价格"
    book.active.append(["产品", "价格"])
    book.active.append(["部件 A", 120])
    book.save(sheet_file)
    sheet_sections = extract_file(sheet_file)
    assert "产品\t价格\n部件 A\t120" in sheet_sections[1].text
    assert "sheet 价格, row 2" == sheet_sections[1].location

    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[6])
    slide.shapes.add_textbox(0, 0, 1000000, 1000000).text = "质保两年"
    slide_file = tmp_path / "deck.pptx"
    slides.save(slide_file)
    assert any("质保两年" in s.text and "slide 1" in s.location for s in extract_file(slide_file))

    html = tmp_path / "page.html"
    html.write_text("<h1>售后说明</h1><script>ignore-me</script><p>支持保修</p>", encoding="utf-8")
    assert "支持保修" in extract_file(html)[0].text
    assert "ignore-me" not in extract_file(html)[0].text

    data = tmp_path / "data.json"
    data.write_text('{"product": "部件 A"}', encoding="utf-8")
    assert "部件 A" in extract_file(data)[0].text


@pytest.mark.skipif(importlib.util.find_spec("reportlab") is None, reason="reportlab test fixture dependency unavailable")
def test_text_pdf_keeps_page_number(tmp_path: Path) -> None:
    from reportlab.pdfgen import canvas

    path = tmp_path / "manual.pdf"
    pdf = canvas.Canvas(str(path))
    pdf.drawString(40, 700, "Warranty period is two years for product A.")
    pdf.save()
    sections = extract_file(path)
    assert sections[0].location == "page 1"
    assert sections[0].method == "text"
    assert "Warranty period" in sections[0].text


def test_large_pdf_page_rejected_before_render(tmp_path: Path) -> None:
    from reportlab.pdfgen import canvas

    path = tmp_path / "huge.pdf"
    pdf = canvas.Canvas(str(path), pagesize=(10000, 10000))
    pdf.showPage()
    pdf.save()
    with pytest.raises(ExtractionError, match="50 megapixel"):
        extract_file(path)


def test_multipage_tiff_keeps_all_pages(tmp_path: Path, monkeypatch) -> None:
    from PIL import Image
    import ragstudio.extract as extraction

    path = tmp_path / "archive.tiff"
    first = Image.new("RGB", (10, 10), "white")
    second = Image.new("RGB", (10, 10), "black")
    first.save(path, save_all=True, append_images=[second])
    monkeypatch.setattr(extraction, "_ocr", lambda image, location, _cmd: [Section(str(image.getpixel((0, 0))), location, "ocr")])
    sections = extract_file(path)
    assert [s.location for s in sections] == ["page 1", "page 2"]
    assert sections[0].text != sections[1].text
